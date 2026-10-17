"""Readable single-column PDF from approved plain text; no invented content."""
from io import BytesIO
from pathlib import Path
import re
import threading
from xml.sax.saxutils import escape
from reportlab import rl_config
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.oxml.ns import qn

lock=threading.Lock()

def pdf_bytes(text,title):
    with lock:
        if 'CareerBody' not in pdfmetrics.getRegisteredFontNames():
            candidates=[Path('C:/Windows/Fonts/arial.ttf'),*(Path(d)/'Vera.ttf' for d in rl_config.TTFSearchPath)]
            font=next((f for f in candidates if f.is_file()),None)
            if font is None: raise RuntimeError('Brak fontu obsługującego polskie znaki.')
            pdfmetrics.registerFont(TTFont('CareerBody',str(font)))
        output=BytesIO(); doc=SimpleDocTemplate(output,pagesize=A4,rightMargin=48,leftMargin=48,topMargin=42,bottomMargin=48,title=title,author='Centrum Kariery')
        heading=ParagraphStyle('CareerHeading',fontName='CareerBody',fontSize=17,leading=23,spaceAfter=12,alignment=TA_LEFT)
        body=ParagraphStyle('CareerBodyStyle',fontName='CareerBody',fontSize=11,leading=16,spaceAfter=8,splitLongWords=True)
        small=ParagraphStyle('CareerSmall',fontName='CareerBody',fontSize=9,leading=13,spaceAfter=12)
        story=[Paragraph(escape(title),heading),Paragraph('SZKIC - do oceny i uzupełnienia przed wysłaniem',small)]
        for line in text.splitlines():
            if line.strip(): story.append(Paragraph(escape(''.join(c for c in line if ord(c)>=32 or c=='\t')),body))
            else: story.append(Spacer(1,5))
        def footer(canvas,document):
            canvas.setFont('CareerBody',9); canvas.drawRightString(A4[0]-48,25,str(document.page))
        doc.build(story,onFirstPage=footer,onLaterPages=footer)
        return output.getvalue()

def docx_bytes(text,title):
    doc=Document(); section=doc.sections[0]
    for style in doc.styles:
        for border in list(style._element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.7)
    normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(8); normal.paragraph_format.line_spacing=1.15
    title_style=doc.styles['Title']; title_style.font.name='Arial'; title_style.font.size=Pt(17); title_style.font.color.rgb=RGBColor(0,0,0); title_style.font.underline=False
    title_style.paragraph_format.space_after=Pt(12)
    doc.add_paragraph(title,'Title')
    note=doc.add_paragraph('SZKIC — do oceny i uzupełnienia przed wysłaniem'); note.runs[0].font.size=Pt(9)
    for line in text.splitlines():
        if line.strip(): doc.add_paragraph(''.join(c for c in line if ord(c)>=32 or c=='\t'))
    doc.core_properties.title=title; doc.core_properties.author='Centrum Kariery'
    output=BytesIO(); doc.save(output); return output.getvalue()

def cv_paragraphs(text):
    return [part.strip() for part in re.split(r'\n\s*\n',text) if part.strip()]

def cv_heading(text):
    from .cv_builder import is_heading
    return is_heading(text) or text in ('Najważniejsze doświadczenie do tej oferty','Materiały portfolio')

def cv_pdf_bytes(text):
    """Export only the approved CV; internal warnings stay in the application."""
    with lock:
        if 'CareerBody' not in pdfmetrics.getRegisteredFontNames():
            font=Path('C:/Windows/Fonts/arial.ttf')
            if not font.is_file(): font=next(Path(d)/'Vera.ttf' for d in rl_config.TTFSearchPath if (Path(d)/'Vera.ttf').is_file())
            pdfmetrics.registerFont(TTFont('CareerBody',str(font)))
        if 'CareerBold' not in pdfmetrics.getRegisteredFontNames():
            font=Path('C:/Windows/Fonts/arialbd.ttf')
            if not font.is_file(): font=next(Path(d)/'VeraBd.ttf' for d in rl_config.TTFSearchPath if (Path(d)/'VeraBd.ttf').is_file())
            pdfmetrics.registerFont(TTFont('CareerBold',str(font)))
        output=BytesIO();doc=SimpleDocTemplate(output,pagesize=A4,rightMargin=48,leftMargin=48,topMargin=42,bottomMargin=42,title='Curriculum vitae',author='')
        name=ParagraphStyle('CVName',fontName='CareerBold',fontSize=19,leading=24,spaceAfter=10,keepWithNext=True)
        heading=ParagraphStyle('CVSection',fontName='CareerBold',fontSize=11.5,leading=16,spaceBefore=12,spaceAfter=6,keepWithNext=True)
        body=ParagraphStyle('CVText',fontName='CareerBody',fontSize=10.5,leading=14,spaceAfter=5,splitLongWords=True)
        story=[]
        for number,part in enumerate(cv_paragraphs(text)):
            story.append(Paragraph(escape(part).replace('\n','<br/>'),name if number==0 else heading if cv_heading(part) else body))
        def footer(canvas,document):
            canvas.setFont('CareerBody',8);canvas.drawRightString(A4[0]-48,22,str(document.page))
        doc.build(story,onFirstPage=footer,onLaterPages=footer);return output.getvalue()

def cv_docx_bytes(text):
    doc=Document();section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=Mm(15)
    section.left_margin=section.right_margin=Mm(17)
    for key,size in [('Normal',10.5),('Title',19),('Heading 2',11.5)]:
        style=doc.styles[key];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        for element in list(style._element.iter(qn('w:pBdr'))):element.getparent().remove(element)
        style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=1.15
    doc.styles['Title'].paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(12)
    for number,part in enumerate(cv_paragraphs(text)):
        doc.add_paragraph(part,'Title' if number==0 else 'Heading 2' if cv_heading(part) else 'Normal')
    doc.core_properties.title='Curriculum vitae';doc.core_properties.author=''
    output=BytesIO();doc.save(output);return output.getvalue()
